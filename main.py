import openpyxl
import re
import os

# mubeen


# Days included in the scheduler
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

# Minimum overlap required for a useful Ehsas Hours match
MIN_OVERLAP = 10


# Convert a time such as 9:00 AM into minutes after midnight
def convert_time(time_string):
    time_string = time_string.strip().upper()

    match = re.fullmatch(
        r"(\d{1,2}):(\d{2})\s*(AM|PM)", time_string
    )

    if not match:
        raise ValueError("Use the format 9:00 AM or 2:30 PM.")

    hour = int(match.group(1))
    minute = int(match.group(2))
    period = match.group(3)

    if hour < 1 or hour > 12 or minute < 0 or minute > 59:
        raise ValueError("Please enter a valid time.")

    if period == "AM":
        hour = hour % 12
    else:
        hour = hour % 12 + 12

    return hour * 60 + minute


# Convert minutes after midnight into a readable time
def format_time(total_minutes):
    hour = total_minutes // 60
    minute = total_minutes % 60

    period = "AM" if hour < 12 else "PM"
    hour = hour % 12

    if hour == 0:
        hour = 12

    return f"{hour}:{minute:02d} {period}"


# Ask the student for a valid time
def get_valid_time(message):
    while True:
        time_string = input(message)

        try:
            return convert_time(time_string)
        except ValueError as error:
            print(error)


# Collect one free-time slot from the student
def get_one_free_slot(day, slot_number):
    print(f"\n{day} - Free-time slot {slot_number}")

    while True:
        start = get_valid_time("Enter start time (e.g. 9:00 AM): ")
        end = get_valid_time("Enter end time (e.g. 10:30 AM): ")

        if start >= end:
            print("End time must be later than start time.")
        else:
            return {
                "day": day,
                "start": start,
                "end": end
            }


# Recursively collect the student's free-time slots
def collect_free_slots(day, number_of_slots, slot_number=1):
    if slot_number > number_of_slots:
        return []

    current_slot = get_one_free_slot(day, slot_number)

    remaining_slots = collect_free_slots(
        day, number_of_slots, slot_number + 1
    )

    return [current_slot] + remaining_slots


# Collect free time for each weekday
def get_student_free_time():
    student_free_slots = []

    print("\n===== ENTER YOUR FREE TIME =====")
    print("Enter the times when you have no classes.")
    print("Use AM/PM, for example 10:00 AM or 2:30 PM.")

    for day in DAYS:
        while True:
            try:
                number_of_slots = int(
                    input(f"\nHow many free-time slots on {day}? ")
                )

                if number_of_slots < 0:
                    print("Please enter zero or a positive number.")
                elif number_of_slots > 20:
                    print("Please enter no more than 20 slots per day.")
                else:
                    break

            except ValueError:
                print("Please enter a whole number.")

        daily_slots = collect_free_slots(day, number_of_slots)
        student_free_slots.extend(daily_slots)

    return student_free_slots


# Interpret a time written in the Excel schedule
def parse_excel_time(time_text, is_end=False, start_minutes=None):
    time_text = str(time_text).strip().upper()
    time_text = time_text.replace(".", ":")
    time_text = re.sub(r"\s+", " ", time_text)

    match = re.fullmatch(
        r"(\d{1,2})(?:\s*:\s*(\d{1,2}))?\s*(AM|PM)?",
        time_text
    )

    if not match:
        return None

    hour = int(match.group(1))
    minute = int(match.group(2) or 0)
    period = match.group(3)

    if hour < 1 or hour > 12 or minute > 59:
        return None

    # Use the stated AM/PM when the spreadsheet provides it.
    if period == "AM":
        hour = hour % 12
    elif period == "PM":
        hour = hour % 12 + 12
    else:
        # The spreadsheet usually omits AM/PM.
        # Assume hours 1-7 mean afternoon/evening.
        if 1 <= hour <= 7:
            hour += 12
        elif hour == 12:
            hour = 12

    total_minutes = hour * 60 + minute

    # If an end time is not later than the start, try the other half
    # of the day. This handles entries such as 11:30 to 1:30.
    if is_end and start_minutes is not None:
        if total_minutes <= start_minutes:
            total_minutes += 12 * 60

        if total_minutes > 24 * 60:
            return None

    return total_minutes


# Read one Excel cell that may contain one or more time ranges
def parse_excel_ranges(cell_value):
    if cell_value is None:
        return []

    text_value = str(cell_value).strip().upper()

    if not text_value or text_value == "-":
        return []

    # Convert decimal-style times: 9.50 becomes 9:50.
    text_value = re.sub(
        r"(?<=\d)\.(?=\d)",
        ":",
        text_value
    )

    # Fix entries such as 11:40:12:40.
    text_value = re.sub(
        r"(\d{1,2}:\d{2}):(\d{1,2}:\d{2})",
        r"\1 - \2",
        text_value
    )

    # Convert compact times such as 500 to 5:00
    # and 1030 to 10:30 when they follow a range separator.
    text_value = re.sub(
        r"(-|TO)\s*(\d{3,4})(?!\d)",
        lambda match: (
            match.group(1) + " " +
            match.group(2)[:-2] + ":" +
            match.group(2)[-2:]
        ),
        text_value
    )

    # Convert formats such as 2 15 to 2:15.
    text_value = re.sub(
        r"(?<!\d)(\d{1,2})\s+(\d{2})(?!\d)",
        r"\1:\2",
        text_value
    )

    time_pattern = (
        r"(\d{1,2}(?:\s*:\s*\d{1,2})?\s*(?:AM|PM)?)"
        r"\s*(?:-|TO)\s*"
        r"(\d{1,2}(?:\s*:\s*\d{1,2})?\s*(?:AM|PM)?)"
    )

    ranges = []

    for match in re.finditer(time_pattern, text_value):
        start_text = match.group(1).strip()
        end_text = match.group(2).strip()

        start = parse_excel_time(start_text)

        if start is None:
            continue

        end = parse_excel_time(
            end_text,
            is_end=True,
            start_minutes=start
        )

        if end is None or end <= start or end > 1440:
            continue

        ranges.append((start, end))

    return ranges

# Import TA availability from the Excel workbook
def get_ehsas_schedule_from_excel():
    filename = "Fall 2026 Schedule of EHSAS Hours.xlsx"
    sheet_name = "Fall 2026 Tutoring Hours"

    if not os.path.exists(filename):
        print("\nERROR: Excel file not found.")
        print("Make sure the Excel file is in the same folder as main.py.")
        return []

    try:
        workbook = openpyxl.load_workbook(
            filename,
            data_only=True,
            read_only=True
        )
    except Exception as error:
        print("\nCould not open the Excel file:", error)
        return []

    if sheet_name not in workbook.sheetnames:
        workbook.close()
        print(f"\nERROR: Worksheet '{sheet_name}' was not found.")
        return []

    sheet = workbook[sheet_name]
    ehsas_sessions = []
    current_course = ""

    # Columns G-K contain Monday-Friday availability.
    for row_number in range(6, sheet.max_row + 1):
        course = sheet.cell(row_number, 3).value
        tutor = sheet.cell(row_number, 4).value

        # Course names may appear only on the first row of a group.
        if course:
            current_course = str(course).strip()

        if not tutor or not current_course:
            continue

        tutor = str(tutor).strip()

        for column_number, day in enumerate(DAYS, start=7):
            cell_value = sheet.cell(row_number, column_number).value

            if not cell_value:
                continue

            time_ranges = parse_excel_ranges(cell_value)

            for start, end in time_ranges:
                session = {
                    "day": day,
                    "subject": current_course,
                    "tutor": tutor,
                    "start": start,
                    "end": end
                }

                ehsas_sessions.append(session)

            if not time_ranges:
                print(
                    "Warning: Could not read Excel row",
                    row_number,
                    "for",
                    day,
                    ":",
                    cell_value
                )

            workbook.close()
            return ehsas_sessions


# Find sessions that overlap with the student's free time
def find_matching_sessions(student_free_slots, ehsas_sessions):
    matches = []

    for free_slot in student_free_slots:
        for session in ehsas_sessions:
            # Sessions on different days cannot match.
            if free_slot["day"] != session["day"]:
                continue

            overlap_start = max(
                free_slot["start"],
                session["start"]
            )

            overlap_end = min(
                free_slot["end"],
                session["end"]
            )

            overlap_minutes = overlap_end - overlap_start

            if overlap_minutes >= MIN_OVERLAP:
                match = {
                    "day": session["day"],
                    "subject": session["subject"],
                    "tutor": session["tutor"],
                    "start": overlap_start,
                    "end": overlap_end,
                    "overlap_minutes": overlap_minutes
                }

                matches.append(match)

    return matches


# Display matching sessions clearly
def display_matches(matches):
    print("\n===== YOUR EHSAS HOURS MATCHES =====")

    if not matches:
        print("No matching sessions were found.")
        print("Try entering more free-time slots.")
        return

    # Sort matches by weekday and then by starting time.
    day_order = {
        day: index for index, day in enumerate(DAYS)
    }

    matches.sort(
        key=lambda match: (
            day_order[match["day"]],
            match["start"],
            match["subject"]
        )
    )

    for number, match in enumerate(matches, start=1):
        print(f"\nMatch {number}")
        print("Day:", match["day"])
        print("Subject:", match["subject"])
        print("Tutor:", match["tutor"])
        print(
            "Available overlap:",
            format_time(match["start"]),
            "to",
            format_time(match["end"])
        )
        print("Overlap duration:", match["overlap_minutes"], "minutes")


# Main function that runs the entire program
def main():
    print("=" * 45)
    print("       EHSAS HOURS SCHEDULER")
    print("=" * 45)

    ehsas_sessions = get_ehsas_schedule_from_excel()

    if not ehsas_sessions:
        print("\nNo Ehsas Hours sessions were imported.")
        print("Check the Excel filename, worksheet, and time formats.")
        return

    print(
        "\nSuccessfully imported",
        len(ehsas_sessions),
        "Ehsas Hours sessions."
    )

    student_free_slots = get_student_free_time()

    if not student_free_slots:
        print("\nYou did not enter any free-time slots.")
        print("There are no matches to display.")
        return

    matches = find_matching_sessions(
        student_free_slots,
        ehsas_sessions
    )

    display_matches(matches)


# Start the program
if __name__ == "__main__":
    main()